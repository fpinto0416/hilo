'''
Compara a versao DIARIA com a versao 4H do HiLo.

DUAS PERGUNTAS DIFERENTES:
  (A) dentro de cada timeframe, long-only bate long+short?
      -> teste pareado sobre a serie da diferenca. Amostra identica, limpo.
  (B) qual timeframe e melhor?
      -> so faz sentido no PERIODO COMUM. A base diaria tem ~26 anos e a de
         4h tem 2-4; comparar os dois inteiros compara quais anos cada um
         pegou, nao os timeframes.

COMO A COMPARACAO E FEITA:
  - os retornos de 4h sao COMPOSTOS DENTRO DO DIA ((1+r).prod()-1), virando
    uma serie diaria. Depois disso tudo anualiza por 252 e as duas viram
    diretamente comparaveis (vol, Sharpe, drawdown, janelas).
  - o periodo e cortado na intersecao das datas.

RESSALVA QUE NAO SOME: no inicio do periodo comum, a estrategia diaria ja
carrega posicao formada por anos de historico anterior, enquanto a de 4h
tambem — mas cada uma com o proprio aquecimento. Diferencas nas primeiras
semanas do recorte refletem estado inicial, nao qualidade do sinal.
'''
import os, sys, numpy as np, pandas as pd
from scipy import stats

# #########################################################################
#                       >>>  CONFIGURACAO  <<<
# Edite aqui e rode com F5 / runfile(). Argumentos de linha de comando, se
# houver, sobrepoem estes valores.
# #########################################################################

ARQ_DIARIO = "saidas_wf_hilo.xlsx"      # saida da versao diaria
ARQ_4H     = "saidas_hilo_4h.xlsx"      # saida da versao 4h

# Os dois hilos precisam ser EQUIVALENTES EM TEMPO DE CALENDARIO.
# hilo_4h = hilo_diario x barras_por_pregao (o hilo_4h.py imprime esse numero
# no arranque: "barras/pregao detectadas"). Com 2 barras/pregao, 50 -> 100.
HILO_DIARIO = 50
HILO_4H     = 100

ARQ_SAIDA  = "comparacao_timeframes.xlsx"

# #########################################################################

ESTRATEGIAS = ["LongShort", "LongOnly", "LongOnly+CDI", "BuyAndHold"]
JANELAS = [30, 60, 90, 180, 250]


def carregar(caminho, hilo=None, agregar_diario=False):
    d = pd.read_excel(caminho, sheet_name="LongOnly_Diario")
    if hilo is not None:
        d = d[d["HiLo_Fixo"] == hilo]
    if d.empty:
        raise ValueError(f"sem dados para hilo={hilo} em {caminho}. "
                         f"hilos disponiveis: {sorted(pd.read_excel(caminho, sheet_name='LongOnly_Diario')['HiLo_Fixo'].unique())}")
    d = d.set_index(pd.DatetimeIndex(d["Data"])).sort_index()
    r = d[[f"Ret_{e}" for e in ESTRATEGIAS]].copy()
    r.columns = ESTRATEGIAS
    if agregar_diario:
        # compoe as barras intradiarias dentro do dia
        r = (1.0 + r).groupby(r.index.normalize()).prod() - 1.0
    else:
        r.index = r.index.normalize()
    return r


def metricas(ret, rot):
    c = (1 + ret).cumprod()
    n = len(ret)
    sd = ret.std(ddof=1)
    out = {"Serie": rot, "N_Dias": n, "Anos": n / 252,
           "Rent_AA": (c.iloc[-1] ** (252 / n) - 1) * 100,
           "Vol_AA": sd * np.sqrt(252) * 100,
           "Sharpe_rf0": ret.mean() / sd * np.sqrt(252) if sd > 0 else np.nan,
           "Drawdown": (c / c.cummax() - 1).min() * 100}
    for w in JANELAS:
        if n > w:
            rr = (c / c.shift(w) - 1).dropna()
            out[f"Pos_{w}d"] = (rr > 0).mean() * 100
    return out


def _nw(x):
    x = np.asarray(x, float); n = len(x)
    l = int(4 * (n / 100) ** (2 / 9))
    e = x - x.mean(); s2 = (e @ e) / n
    for k in range(1, l + 1):
        s2 += 2 * (1 - k / (l + 1)) * ((e[k:] @ e[:-k]) / n)
    t = x.mean() / np.sqrt(s2 / n)
    return t, 2 * (1 - stats.norm.cdf(abs(t)))


def pareado(a, b, rot):
    d = (np.asarray(a, float) - np.asarray(b, float))
    d = d[np.isfinite(d)]
    n = len(d)
    if n < 60 or d.std(ddof=1) == 0:
        return None
    t = d.mean() / (d.std(ddof=1) / np.sqrt(n))
    tn, pn = _nw(d)
    return {"Comparacao": rot, "N_Dias": n,
            "Dif_AA_pp": d.mean() * 252 * 100,
            "t": t, "p_Padrao": 2 * (1 - stats.norm.cdf(abs(t))),
            "p_NeweyWest": pn}


def mcnemar(ca, cb, rot, w):
    n = len(ca)
    if n <= w:
        return None
    ct = np.arange(0, n, w)
    if len(ct) < 4:
        return None
    pa = (ca[ct[1:]] / ca[ct[:-1]] - 1) > 0
    pb = (cb[ct[1:]] / cb[ct[:-1]] - 1) > 0
    b_ = int((pa & ~pb).sum()); c_ = int((~pa & pb).sum())
    disc = b_ + c_
    return {"Comparacao": rot, "Janela_d": w, "N_Janelas": len(ct) - 1,
            "Pos_A": pa.mean() * 100, "Pos_B": pb.mean() * 100,
            "So_A": b_, "So_B": c_, "Discordantes": disc,
            "p": stats.binomtest(b_, disc, 0.5).pvalue if disc else np.nan}


def _barras_por_pregao(caminho, hilo):
    """Mede barras/pregao no proprio arquivo de 4h, para checar a equivalencia."""
    d = pd.read_excel(caminho, sheet_name="LongOnly_Diario")
    d = d[d["HiLo_Fixo"] == hilo]
    idx = pd.DatetimeIndex(d["Data"])
    por_dia = pd.Series(1, index=idx).groupby(idx.normalize()).sum()
    return float(por_dia.median())


def main(arq_diario, arq_4h, hilo_diario, hilo_4h, saida="comparacao_timeframes.xlsx"):
    D = carregar(arq_diario, hilo_diario, agregar_diario=False)
    H = carregar(arq_4h, hilo_4h, agregar_diario=True)

    # CHECAGEM DE EQUIVALENCIA: comparar lookbacks de tamanhos diferentes e o
    # erro mais facil de cometer aqui, e ele nao gera nenhum sintoma — o script
    # roda e devolve numeros que parecem certos.
    bpp = _barras_por_pregao(arq_4h, hilo_4h)
    pregoes_4h = hilo_4h / bpp
    print(f"\nEQUIVALENCIA: {bpp:.2f} barras/pregao no arquivo de 4h")
    print(f"  hilo diario {hilo_diario} = {hilo_diario} pregoes de lookback")
    print(f"  hilo 4h {hilo_4h} = {pregoes_4h:.0f} pregoes de lookback")
    desvio = abs(pregoes_4h - hilo_diario) / hilo_diario
    if desvio > 0.15:
        print(f"  *** AVISO: lookbacks diferem {desvio:.0%}. Nao e comparacao de "
              f"timeframe, e de janela.\n"
              f"      Para equivaler a {hilo_diario} pregoes, use HILO_4H = "
              f"{int(round(hilo_diario * bpp))} (e rode o hilo_4h.py com esse hilo fixo).")
    else:
        print(f"  ok (desvio de {desvio:.0%})")
    com = D.index.intersection(H.index)
    print(f"diario : {D.index.min().date()} a {D.index.max().date()} ({len(D)} dias)")
    print(f"4h     : {H.index.min().date()} a {H.index.max().date()} ({len(H)} dias apos agregacao)")
    print(f"COMUM  : {com.min().date()} a {com.max().date()} ({len(com)} dias, {len(com)/252:.1f} anos)")
    if len(com) < 250:
        print("AVISO: periodo comum abaixo de 1 ano. Qualquer conclusao aqui e ruido.")
    D, H = D.loc[com], H.loc[com]

    linhas = [metricas(D[e], f"Diario h{hilo_diario} | {e}") for e in ESTRATEGIAS]
    linhas += [metricas(H[e], f"4h h{hilo_4h} | {e}") for e in ESTRATEGIAS]
    met = pd.DataFrame(linhas)

    testes = []
    # (A) dentro de cada timeframe
    for rot, F in [("Diario", D), ("4h", H)]:
        for A, B in [("LongOnly+CDI", "LongShort"), ("LongOnly", "LongShort"),
                     ("LongOnly+CDI", "BuyAndHold")]:
            t = pareado(F[A], F[B], f"[{rot}] {A} - {B}")
            if t: testes.append(t)
    # (B) entre timeframes, mesma estrategia
    for e in ESTRATEGIAS:
        t = pareado(H[e], D[e], f"[4h - Diario] {e}")
        if t: testes.append(t)
    tst = pd.DataFrame(testes)

    mc = []
    for e in ESTRATEGIAS:
        ca = (1 + H[e]).cumprod().values; cb = (1 + D[e]).cumprod().values
        for w in JANELAS:
            m = mcnemar(ca, cb, f"[4h - Diario] {e}", w)
            if m: mc.append(m)
    for rot, F in [("Diario", D), ("4h", H)]:
        ca = (1 + F["LongOnly+CDI"]).cumprod().values
        cb = (1 + F["LongShort"]).cumprod().values
        for w in JANELAS:
            m = mcnemar(ca, cb, f"[{rot}] LongOnly+CDI - LongShort", w)
            if m: mc.append(m)
    mcn = pd.DataFrame(mc)

    with pd.ExcelWriter(saida, engine="openpyxl") as wr:
        met.to_excel(wr, index=False, sheet_name="Metricas")
        tst.to_excel(wr, index=False, sheet_name="Testes_Pareados")
        mcn.to_excel(wr, index=False, sheet_name="McNemar_Janelas")
    print(f"\nsalvo: {saida}")
    pd.set_option("display.width", 320)
    print("\n=== METRICAS (periodo comum) ===")
    print(met.round(2).to_string(index=False))
    print("\n=== TESTES PAREADOS ===")
    print(tst.round(4).to_string(index=False))
    return met, tst, mcn


def _checar(caminho, rotulo):
    """Falha com mensagem util em vez de stack trace do openpyxl."""
    if not os.path.exists(caminho):
        raise FileNotFoundError(
            f"{rotulo} nao encontrado: {os.path.abspath(caminho)}\n"
            f"  diretorio de trabalho: {os.path.abspath(os.getcwd())}\n"
            "  ajuste ARQ_DIARIO / ARQ_4H no topo do arquivo, ou use caminho absoluto."
        )
    try:
        from openpyxl import load_workbook
        abas = load_workbook(caminho, read_only=True).sheetnames
    except Exception as e:
        raise ValueError(f"nao foi possivel abrir {rotulo}: {e}")
    if "LongOnly_Diario" not in abas:
        raise ValueError(
            f"{rotulo} nao tem a aba 'LongOnly_Diario'.\n"
            f"  abas presentes: {abas}\n"
            "  Isso costuma indicar rodada INCOMPLETA (o script quebrou antes de\n"
            "  gravar as abas de long-only). Rode o hilo novamente ate o fim."
        )


if __name__ == "__main__":
    # linha de comando sobrepoe o painel, se fornecida
    if len(sys.argv) >= 5:
        ARQ_DIARIO, ARQ_4H = sys.argv[1], sys.argv[2]
        HILO_DIARIO, HILO_4H = int(sys.argv[3]), int(sys.argv[4])

    print(f"diario: {ARQ_DIARIO} (hilo {HILO_DIARIO})")
    print(f"4h    : {ARQ_4H} (hilo {HILO_4H})")
    _checar(ARQ_DIARIO, "arquivo diario")
    _checar(ARQ_4H, "arquivo 4h")
    main(ARQ_DIARIO, ARQ_4H, HILO_DIARIO, HILO_4H, ARQ_SAIDA)
